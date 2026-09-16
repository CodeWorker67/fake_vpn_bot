import asyncio
import logging
import os
import tempfile
from datetime import datetime

import openpyxl
from aiogram import Router
from aiogram.filters import Command
from aiogram.types import FSInputFile, Message
from openpyxl.styles import Alignment, Border, Side

from config import ADMIN_IDS
from stats_db import BotUserRow, count_users, list_all_users

logger = logging.getLogger(__name__)
router = Router()

_EXCEL_COL_WIDTH_MAX = 255

_EXPORT_COLUMNS = (
    ("user_id", "User ID"),
    ("username", "Username"),
    ("first_name", "First name"),
    ("last_name", "Last name"),
    ("is_premium", "Premium"),
    ("joined_at", "Joined at"),
)


def _excel_scalar(value):
    if value is None:
        return value
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, bool):
        return "yes" if value else "no"
    return value


def _sync_build_users_export(users: list[BotUserRow]) -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "users"

    header_alignment = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    for col_num, (_, title) in enumerate(_EXPORT_COLUMNS, 1):
        cell = ws.cell(row=1, column=col_num, value=title)
        cell.alignment = header_alignment
        cell.border = thin_border

    for row_num, user in enumerate(users, 2):
        for col_num, (attr, _) in enumerate(_EXPORT_COLUMNS, 1):
            cell = ws.cell(
                row=row_num,
                column=col_num,
                value=_excel_scalar(getattr(user, attr)),
            )
            cell.border = thin_border

    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_len + 2, _EXCEL_COL_WIDTH_MAX)

    ws.freeze_panes = ws["A2"]

    fd, path = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    wb.save(path)
    return path


@router.message(Command(commands=["export"]))
async def export_users_to_excel(message: Message) -> None:
    if not message.from_user or message.from_user.id not in ADMIN_IDS:
        await message.answer("❌ Эта команда доступна только администраторам.")
        return

    try:
        total = await asyncio.to_thread(count_users)
        await message.answer(f"🔄 В базе {total} пользователей. Формирую Excel…")

        users = await asyncio.to_thread(list_all_users)
        export_path = await asyncio.to_thread(_sync_build_users_export, users)

        now_s = datetime.now().strftime("%d.%m.%Y %H:%M")
        caption = (
            "📊 Экспорт пользователей бота\n"
            f"📅 Создано: {now_s}\n\n"
            f"👥 Записей: {len(users)}"
        )
        try:
            await message.answer_document(
                document=FSInputFile(export_path, filename="bot_users_export.xlsx"),
                caption=caption,
            )
        finally:
            try:
                os.remove(export_path)
            except OSError:
                pass

        logger.info("Администратор %s экспортировал пользователей в Excel", message.from_user.id)
    except Exception as e:
        logger.exception("Ошибка экспорта пользователей")
        await message.answer(f"❌ Ошибка при экспорте: {e}")
